"""Loaders for the paper's small benchmark datasets, with the original preprocessing.

The raw files live outside the repository, in ``$COSTAL_DATA`` (default
``/Volumes/X31/costal-plus-data``), one directory per dataset; ``SOURCES`` says where each came
from. Parsed matrices are cached as ``.npz`` under ``$COSTAL_DATA/cache``.

Labels (``-1`` = unlabeled):

- **Levine_13** (mass cytometry, 13 markers): the 24 manually gated populations; arcsinh(x/5).
- **E-MTAB-3321** (Goolam 2016, Smart-seq2, mouse embryo): developmental stage from the SDRF
  (2-, 4-, 8-, 16-, 32-cell), the "known cell source" labels of the paper.
- **GSE81861** (Li 2017, cell lines): the cell line encoded in the column name, with the batch
  suffix (``_B1`` / ``_B2``) dropped — 7 lines.
- **GSE74672** (Romanov 2017, hypothalamus): ``level1 class`` (7 classes).
- **GSE84133** (Baron 2016, pancreas, donor human3): ``assigned_cluster`` (14 types).
- **Samusik_01** (mass cytometry, mouse bone marrow): 24 gated populations; the 39 ``type`` markers;
  arcsinh(x/5).
- **Nilsson_rare** (flow cytometry): HSCs (358 cells) vs the rest; the 13 ``type`` markers plus PI;
  arcsinh(x/150).
- **Mosmann_rare** (flow cytometry): activated cells (109) vs the rest; the 7 ``type`` and 7 ``state``
  markers plus Live_Dead; arcsinh(x/150).

The three cytometry sets come from Bioconductor's HDCytoData (ExperimentHub, the ``*_SE.rda``
SummarizedExperiments), read with ``bench/rds.py`` instead of R. Their marker sets are the ones that
reproduce the paper's Table 2 exactly; Table 2's "non-zeros" counts *positive* entries (docs/03 §10.1).

``preprocess`` reproduces the original ``pp()`` (bench/reference, ``CosTaL/clustering.py``) and
tracks which cells survive ``filter_cells``, which the original does not (it returns only the
matrix, so labels would silently misalign if a cell were dropped). ``check`` compares the result
with Table 2 of the paper.

Usage: python bench/datasets.py check
"""

from __future__ import annotations

import gzip
import io
import json
import os
import sys
import warnings
from dataclasses import dataclass
from pathlib import Path

import numpy as np
import scipy.sparse as sp

HERE = Path(__file__).resolve().parent
if str(HERE) not in sys.path:
    sys.path.insert(0, str(HERE))  # for bench/rds.py
DATA = Path(os.environ.get("COSTAL_DATA", "/Volumes/X31/costal-plus-data"))
CACHE = DATA / "cache"

SOURCES = {
    "Levine_13": "FlowRepository FR-FCM-ZZPH via the Weber & Robinson 2016 benchmark: "
                 "Levine_13dim_notransform.fcs, population_names_Levine_13dim.txt",
    "E-MTAB-3321": "ArrayExpress E-MTAB-3321: Goolam_et_al_2015_count_table.tsv, E-MTAB-3321.sdrf.txt",
    "GSE81861": "GEO GSE81861: GSE81861_Cell_Line_COUNT.csv.gz",
    "GSE74672": "GEO GSE74672: GSE74672_expressed_mols_with_classes.xlsx.gz",
    "GSE84133": "GEO GSM2230759 (GSE84133, donor human3): GSM2230759_human3_umifm_counts.csv.gz",
    "Samusik_01": "Bioconductor ExperimentHub EH2244 (HDCytoData): Samusik_01_SE.rda",
    "Nilsson_rare": "Bioconductor ExperimentHub EH2248 (HDCytoData): Nilsson_rare_SE.rda",
    "Mosmann_rare": "Bioconductor ExperimentHub EH2250 (HDCytoData): Mosmann_rare_SE.rda",
}

# Table 2 of the paper: cells after preprocessing, features, non-zeros, sparsity (%)
PAPER_TABLE2 = {
    "Levine_13": {"cells": 167_044, "labeled": 81_747, "features": 13, "nnz": 1_689_738, "sparsity": 22.188},
    "E-MTAB-3321": {"original_features": 41_480, "cells": 124, "features": 2000, "nnz": 144_499, "sparsity": 41.734},
    "GSE81861": {"original_features": 57_241, "cells": 561, "features": 2000, "nnz": 214_688, "sparsity": 80.866},
    "GSE74672": {"original_features": 24_341, "cells": 2_881, "features": 2000, "nnz": 821_520, "sparsity": 85.742},
    "GSE84133": {"original_features": 20_125, "cells": 3_605, "features": 2000, "nnz": 792_390, "sparsity": 89.010},
    "Samusik_01": {"cells": 86_864, "labeled": 53_173, "features": 39, "nnz": 1_949_437, "sparsity": 42.455},
    "Nilsson_rare": {"cells": 44_140, "labeled": 358, "features": 14, "nnz": 616_569, "sparsity": 0.225},
    "Mosmann_rare": {"cells": 396_460, "labeled": 109, "features": 15, "nnz": 5_808_969, "sparsity": 2.319},
}
NAMES = ["Levine_13", "E-MTAB-3321", "GSE81861", "GSE74672", "GSE84133"]
SCRNA = [n for n in NAMES if n != "Levine_13"]
CYTOMETRY_EXTRA = ["Samusik_01", "Nilsson_rare", "Mosmann_rare"]
# arcsinh cofactor platform of each cytometry set; the rare population of the one-population sets
PLATFORM = {"Levine_13": "mass", "Samusik_01": "mass", "Nilsson_rare": "flow", "Mosmann_rare": "flow"}
RARE = {"Nilsson_rare": "HSCs", "Mosmann_rare": "activated"}


@dataclass
class Dataset:
    name: str
    kind: str  # "cytometry" or "scrnaseq"
    x: np.ndarray | sp.csr_matrix  # cytometry: raw intensities; scRNA-seq: raw counts, cells x features
    labels: np.ndarray  # int64, -1 = unlabeled
    classes: list[str]
    features: list[str]


def _encode(names) -> tuple[np.ndarray, list[str]]:
    classes = sorted({str(s) for s in names if s is not None and str(s) != "nan"})
    lookup = {c: i for i, c in enumerate(classes)}
    labels = np.array([lookup.get(str(s), -1) for s in names], dtype=np.int64)
    return labels, classes


def _cached(name: str, parse) -> Dataset:
    path = CACHE / f"{name}.npz"
    if path.exists():
        z = np.load(path, allow_pickle=False)
        meta = json.loads(str(z["meta"]))
        x = (sp.csr_matrix((z["data"], z["indices"], z["indptr"]), shape=tuple(z["shape"]))
             if meta["sparse"] else z["x"])
        return Dataset(name, meta["kind"], x, z["labels"], meta["classes"], meta["features"])
    ds = parse()
    CACHE.mkdir(parents=True, exist_ok=True)
    meta = json.dumps({"kind": ds.kind, "classes": ds.classes, "features": ds.features,
                       "sparse": sp.issparse(ds.x)})
    if sp.issparse(ds.x):
        np.savez(path, data=ds.x.data, indices=ds.x.indices, indptr=ds.x.indptr,
                 shape=np.array(ds.x.shape), labels=ds.labels, meta=meta)
    else:
        np.savez(path, x=ds.x, labels=ds.labels, meta=meta)
    return ds


def _levine_13() -> Dataset:
    import fcsparser

    d = DATA / "Levine_13"
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        _, df = fcsparser.parse(str(d / "Levine_13dim_notransform.fcs"), reformat_meta=True)
    markers = [c for c in df.columns if c != "label"]
    names = {}
    for line in (d / "population_names_Levine_13dim.txt").read_text().splitlines()[1:]:
        k, v = line.split("\t")
        names[float(k)] = v
    lab = df["label"].to_numpy()
    labels, classes = _encode([names.get(v) for v in lab])
    return Dataset("Levine_13", "cytometry", np.ascontiguousarray(df[markers].to_numpy(np.float64)),
                   labels, classes, markers)


def _goolam() -> Dataset:
    import pandas as pd

    d = DATA / "E-MTAB-3321"
    counts = pd.read_csv(d / "Goolam_et_al_2015_count_table.tsv", sep="\t", index_col=0)
    sdrf = pd.read_csv(d / "E-MTAB-3321.sdrf.txt", sep="\t").drop_duplicates("Source Name")
    stage = dict(zip(sdrf["Source Name"], sdrf["Characteristics[developmental stage]"]))
    labels, classes = _encode([stage[c].replace("cleavage ", "") for c in counts.columns])
    return Dataset("E-MTAB-3321", "scrnaseq", sp.csr_matrix(counts.to_numpy(np.float32).T),
                   labels, classes, counts.index.tolist())


def _li() -> Dataset:
    import pandas as pd

    counts = pd.read_csv(DATA / "GSE81861" / "GSE81861_Cell_Line_COUNT.csv.gz", index_col=0)
    lines = [c.split("__")[1].split("_B")[0] for c in counts.columns]
    labels, classes = _encode(lines)
    return Dataset("GSE81861", "scrnaseq", sp.csr_matrix(counts.to_numpy(np.float32).T),
                   labels, classes, counts.index.tolist())


def _romanov() -> Dataset:
    import openpyxl

    raw = gzip.open(DATA / "GSE74672" / "GSE74672_expressed_mols_with_classes.xlsx.gz").read()
    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    rows = wb[wb.sheetnames[0]].iter_rows(values_only=True)
    header = [next(rows) for _ in range(12)]
    assert header[0][0] == "cellID" and header[1][0] == "level1 class", header[:2]
    genes, cols = [], []
    for row in rows:
        genes.append(str(row[0]))
        cols.append(np.asarray(row[1:], dtype=np.float32))
    x = sp.csr_matrix(np.vstack(cols).T)
    labels, classes = _encode(header[1][1:])
    return Dataset("GSE74672", "scrnaseq", x, labels, classes, genes)


def _baron() -> Dataset:
    import pandas as pd

    df = pd.read_csv(DATA / "GSE84133" / "GSM2230759_human3_umifm_counts.csv.gz", index_col=0)
    labels, classes = _encode(df["assigned_cluster"])
    genes = [c for c in df.columns if c not in ("barcode", "assigned_cluster")]
    return Dataset("GSE84133", "scrnaseq", sp.csr_matrix(df[genes].to_numpy(np.float32)),
                   labels, classes, genes)


def _hdcyto(name: str, classes: tuple, extra: tuple, unlabeled: str | None) -> Dataset:
    from rds import read_rda, summarized_experiment

    se = next(iter(read_rda(DATA / name / f"{name}_SE.rda").values()))
    x, rows, cols = summarized_experiment(se)
    keep = [i for i, (m, c) in enumerate(zip(cols["marker_name"], cols["marker_class"]))
            if c in classes or m in extra]
    pops = [None if p == unlabeled else p for p in rows["population_id"]]
    labels, names = _encode(pops)
    return Dataset(name, "cytometry", np.ascontiguousarray(x[:, keep], dtype=np.float64), labels, names,
                   [cols["marker_name"][i] for i in keep])


_PARSERS = {"Levine_13": _levine_13, "E-MTAB-3321": _goolam, "GSE81861": _li,
            "GSE74672": _romanov, "GSE84133": _baron,
            "Samusik_01": lambda: _hdcyto("Samusik_01", ("type",), (), "unassigned"),
            "Nilsson_rare": lambda: _hdcyto("Nilsson_rare", ("type",), ("PI",), None),
            "Mosmann_rare": lambda: _hdcyto("Mosmann_rare", ("type", "state"), ("Live_Dead",), None)}


def load(name: str) -> Dataset:
    return _cached(name, _PARSERS[name])


def preprocess(ds: Dataset, preset: str = "paper") -> tuple[np.ndarray, np.ndarray]:
    """The original ``pp()`` via ``costal.preprocess`` (docs/04 §6.7): the dense float32
    matrix and the indices of the cells kept. Cytometry: arcsinh with the platform's cofactor."""
    import costal as cp

    if ds.kind == "cytometry":
        return cp.preprocess(ds.x, PLATFORM[ds.name], return_cells=True)
    return cp.preprocess(ds.x, "scrnaseq", preset=preset, return_cells=True)


def describe(ds: Dataset, x: np.ndarray, kept: np.ndarray) -> dict:
    labels = ds.labels[kept]
    labeled = (int(np.sum(labels == ds.classes.index(RARE[ds.name]))) if ds.name in RARE
               else int(np.sum(labels >= 0)))  # Table 2 counts the rare population as "labeled"
    positives = int(np.count_nonzero(x > 0))
    return {
        "original_features": ds.x.shape[1], "cells": int(x.shape[0]), "features": int(x.shape[1]),
        "labeled": labeled, "classes": len(ds.classes),
        "nnz": int(np.count_nonzero(x)), "sparsity": round(100 * (1 - np.count_nonzero(x) / x.size), 3),
        "positives": positives, "nonpositive_share": round(100 * (1 - positives / x.size), 3),
        "negative_share": round(100 * float(np.mean(x < 0)), 3),
        "smallest_class": int(np.bincount(labels[labels >= 0]).min()),
    }


# Table 2's "non-zeros" and "sparsity" count positive entries (equal to non-zeros on scRNA-seq)
_TABLE2_KEYS = {"nnz": "positives", "sparsity": "nonpositive_share"}


def cmd_check() -> None:
    for name in PAPER_TABLE2:
        ds = load(name)
        x, kept = preprocess(ds)
        ours = describe(ds, x, kept)
        paper = PAPER_TABLE2[name]
        diff = {k: (ours[_TABLE2_KEYS.get(k, k)], v) for k, v in paper.items()
                if ours.get(_TABLE2_KEYS.get(k, k)) != v}
        print(f"{name:<12} {json.dumps(ours)}")
        print(f"{'':<12} differs from the paper (ours, paper): {diff or 'nothing'}", flush=True)


if __name__ == "__main__":
    {"check": cmd_check}[sys.argv[1] if len(sys.argv) > 1 else "check"]()
